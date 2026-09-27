from io import BytesIO

from docx import Document as DocxDocument
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw

from tests.fakes import InMemoryVectorStore


def create_knowledge_base(client: TestClient, name: str) -> str:
    response = client.post("/api/v1/knowledge", json={"name": name})
    assert response.status_code == 201
    return response.json()["id"]


def test_docx_upload_parses_paragraphs_and_tables(
    client: TestClient,
    vector_store: InMemoryVectorStore,
) -> None:
    document = DocxDocument()
    document.add_heading("Project Notes", level=1)
    document.add_paragraph("618 promotion analysis project")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "AUC 0.85"
    buffer = BytesIO()
    document.save(buffer)
    knowledge_base_id = create_knowledge_base(client, "DOCX Test")

    response = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": knowledge_base_id},
        files={
            "file": (
                "notes.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 202
    stored = client.get("/api/v1/documents").json()[0]
    assert stored["status"] == "ready"
    content = " ".join(record.content for record in vector_store.records.values())
    assert "618 promotion analysis" in content
    assert "AUC 0.85" in content


def test_image_upload_uses_local_ocr(
    client: TestClient,
    vector_store: InMemoryVectorStore,
) -> None:
    image = Image.new("RGB", (800, 180), "white")
    draw = ImageDraw.Draw(image)
    draw.text((30, 60), "RAG IMAGE TEST 618", fill="black")
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    knowledge_base_id = create_knowledge_base(client, "Image Test")

    response = client.post(
        "/api/v1/documents/upload",
        data={"knowledge_base_id": knowledge_base_id},
        files={"file": ("scan.png", buffer.getvalue(), "image/png")},
    )

    assert response.status_code == 202
    stored = client.get("/api/v1/documents").json()[0]
    assert stored["status"] == "ready"
    content = " ".join(record.content for record in vector_store.records.values())
    assert "RAG" in content
    assert "618" in content
