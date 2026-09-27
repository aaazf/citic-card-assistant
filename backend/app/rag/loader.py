import re
import zipfile
from dataclasses import dataclass, field
from functools import lru_cache
from html.parser import HTMLParser
from pathlib import Path
from posixpath import join as posix_join
from posixpath import normpath
from typing import Any, Protocol
from urllib.parse import unquote
from xml.etree import ElementTree

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".pdf",
    ".txt",
    ".md",
    ".markdown",
    ".docx",
    ".epub",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
EPUB_CONTAINER_PATH = "META-INF/container.xml"
EPUB_CONTENT_TYPES = {
    "application/xhtml+xml",
    "text/html",
}
EPUB_BLOCK_TAGS = {
    "address",
    "article",
    "aside",
    "blockquote",
    "br",
    "dd",
    "div",
    "dl",
    "dt",
    "figcaption",
    "figure",
    "footer",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "header",
    "hr",
    "li",
    "main",
    "ol",
    "p",
    "pre",
    "section",
    "table",
    "tbody",
    "td",
    "tfoot",
    "th",
    "thead",
    "tr",
    "ul",
}
EPUB_SKIP_TAGS = {"script", "style", "noscript", "svg", "template"}
EPUB_SKIP_TAGS.update({"head", "nav"})
EPUB_HEADING_TAGS = {"h1", "h2", "h3", "h4", "h5", "h6"}


@dataclass(frozen=True)
class LoadedSection:
    content: str
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class LoadedDocument:
    sections: list[LoadedSection]


class DocumentLoader(Protocol):
    def load(self, path: Path) -> LoadedDocument: ...


class _EpubTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.heading_candidates: list[str] = []
        self._heading_parts: list[str] = []
        self.title_parts: list[str] = []
        self._skip_depth = 0
        self._heading_depth = 0
        self._title_depth = 0

    @staticmethod
    def _local_name(tag: str) -> str:
        return tag.rsplit("}", 1)[-1].lower()

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        name = self._local_name(tag)
        if name in EPUB_SKIP_TAGS:
            self._skip_depth += 1
            return
        if self._skip_depth:
            return
        if name == "br":
            self.parts.append("\n")
        elif name in EPUB_BLOCK_TAGS:
            self.parts.append("\n")
        if name == "title":
            self._title_depth += 1
        if name in EPUB_HEADING_TAGS:
            if self._heading_depth == 0:
                self._heading_parts = []
            self._heading_depth += 1

    def handle_endtag(self, tag: str) -> None:
        name = self._local_name(tag)
        if name in EPUB_SKIP_TAGS:
            self._skip_depth = max(0, self._skip_depth - 1)
            return
        if self._skip_depth:
            return
        if name == "title":
            self._title_depth = max(0, self._title_depth - 1)
        if name in EPUB_HEADING_TAGS:
            self._heading_depth = max(0, self._heading_depth - 1)
            if self._heading_depth == 0:
                title = re.sub(r"\s+", " ", "".join(self._heading_parts)).strip()
                if title:
                    self.heading_candidates.append(title)
        if name in EPUB_BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth or not data:
            return
        if self._title_depth:
            self.title_parts.append(data)
            return
        self.parts.append(data)
        if self._heading_depth:
            self._heading_parts.append(data)

    def text(self) -> str:
        lines = [
            re.sub(r"[ \t\f\v]+", " ", line).strip()
            for line in "".join(self.parts).splitlines()
        ]
        return "\n".join(line for line in lines if line)

    def chapter_title(self) -> str | None:
        candidates = (*self.heading_candidates, "".join(self.title_parts))
        for candidate in candidates:
            title = re.sub(r"\s+", " ", candidate).strip()
            if title:
                return title[:160]
        return None


@lru_cache(maxsize=1)
def _get_ocr_engine():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR()


class LocalDocumentLoader:
    def load(self, path: Path) -> LoadedDocument:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return self._load_pdf(path)
        if suffix == ".docx":
            return self._load_docx(path)
        if suffix == ".epub":
            return self._load_epub(path)
        if suffix in IMAGE_EXTENSIONS:
            return self._load_image(path)
        if suffix in {".txt", ".md", ".markdown"}:
            return self._load_text(path)
        raise ValueError(f"Unsupported document type: {suffix}")

    @staticmethod
    def _load_text(path: Path) -> LoadedDocument:
        content = path.read_text(encoding="utf-8")
        return LoadedDocument(sections=[LoadedSection(content=content)])

    @staticmethod
    def _load_pdf(path: Path) -> LoadedDocument:
        reader = PdfReader(str(path))
        sections = [
            LoadedSection(content=page.extract_text() or "", metadata={"page": index})
            for index, page in enumerate(reader.pages, start=1)
        ]
        return LoadedDocument(sections=sections)

    @staticmethod
    def _load_docx(path: Path) -> LoadedDocument:
        document = DocxDocument(str(path))
        blocks = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    blocks.append(" | ".join(values))
        content = "\n".join(blocks)
        return LoadedDocument(
            sections=[
                LoadedSection(
                    content=content,
                    metadata={"source_type": "docx"},
                )
            ]
        )

    @staticmethod
    def _load_epub(path: Path) -> LoadedDocument:
        with zipfile.ZipFile(path) as archive:
            package_path = LocalDocumentLoader._epub_package_path(archive)
            package_dir = str(Path(package_path).parent).replace("\\", "/")
            package = ElementTree.fromstring(archive.read(package_path))
            manifest = {
                item.get("id"): item
                for item in package.iter()
                if item.tag.rsplit("}", 1)[-1].lower() == "item" and item.get("id")
            }
            spine_ids = [
                item.get("idref")
                for item in package.iter()
                if item.tag.rsplit("}", 1)[-1].lower() == "itemref"
            ]
            hrefs = [
                manifest[item_id].get("href")
                for item_id in spine_ids
                if item_id in manifest
            ]
            if not any(hrefs):
                hrefs = [
                    item.get("href")
                    for item in manifest.values()
                    if item.get("media-type") in EPUB_CONTENT_TYPES
                ]

            sections: list[LoadedSection] = []
            for index, href in enumerate(hrefs, start=1):
                if not href:
                    continue
                member = LocalDocumentLoader._resolve_epub_path(package_dir, href)
                if member not in archive.namelist():
                    continue
                parser = _EpubTextParser()
                parser.feed(archive.read(member).decode("utf-8-sig", errors="replace"))
                parser.close()
                content = parser.text()
                if not content:
                    continue
                section_metadata: dict[str, Any] = {
                    "source_type": "epub",
                    "chapter_index": index,
                }
                chapter_title = parser.chapter_title()
                if chapter_title:
                    section_metadata["chapter_title"] = chapter_title
                sections.append(LoadedSection(content=content, metadata=section_metadata))

        if not sections:
            raise ValueError("No extractable text was found in the EPUB document.")
        return LoadedDocument(sections=sections)

    @staticmethod
    def _epub_package_path(archive: zipfile.ZipFile) -> str:
        try:
            container = ElementTree.fromstring(archive.read(EPUB_CONTAINER_PATH))
        except (KeyError, ElementTree.ParseError) as exc:
            raise ValueError("Invalid EPUB: missing or malformed container.xml.") from exc
        rootfile = next(
            (
                element.get("full-path")
                for element in container.iter()
                if element.tag.rsplit("}", 1)[-1].lower() == "rootfile"
            ),
            None,
        )
        if not rootfile or rootfile not in archive.namelist():
            raise ValueError("Invalid EPUB: package document was not found.")
        return rootfile

    @staticmethod
    def _resolve_epub_path(package_dir: str, href: str) -> str:
        clean_href = unquote(href.split("#", 1)[0]).lstrip("/")
        return normpath(posix_join(package_dir, clean_href))

    @staticmethod
    def _load_image(path: Path) -> LoadedDocument:
        result, _ = _get_ocr_engine()(str(path))
        lines = [str(row[1]).strip() for row in (result or []) if len(row) > 1 and row[1]]
        return LoadedDocument(
            sections=[
                LoadedSection(
                    content="\n".join(lines),
                    metadata={"source_type": "image", "page": 1, "ocr": True},
                )
            ]
        )
